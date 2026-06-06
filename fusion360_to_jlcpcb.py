#!/usr/bin/env python3
"""
Transform Autodesk Fusion 360 PCB assembly exports (BOM + CPL CSVs)
to the format required by JLCPCB (BOM + CPL xlsx).

Usage:
    python3 fusion360_to_jlcpcb.py BOM.csv FRONT_CPL.csv [BACK_CPL.csv] [options]

Arguments:
    BOM.csv         Fusion 360 BOM export (CSV)
    FRONT_CPL.csv   Fusion 360 Pick-and-Place front-side CSV
    BACK_CPL.csv    Fusion 360 Pick-and-Place back-side CSV (omit for single-sided boards)

Options:
    --bom-out PATH  Output path for BOM xlsx (default: ./bom.xlsx)
    --cpl-out PATH  Output path for CPL xlsx (default: ./cpl.xlsx)
    -h, --help      Show this help message
"""

import argparse
import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# BOM helpers
# ---------------------------------------------------------------------------

_genericDescriptions = {
  "capacitor - generic",
  "resistor fixed - generic",
  "transistor",
  "diode rectifier - generic",
  "voltage regulator",
  "n-channel mosfet - generic",
  "diode",
}


def _IsGeneric(description: str) -> bool:
  return description.strip().lower() in _genericDescriptions


def _ChooseComment(value: str, device: str, description: str) -> str:
  """
  Heuristic for the JLCPCB 'Comment' field:
    1. If Value differs from Device, Value is a component value (e.g. '10k') — use it.
    2. Otherwise, if Description is non-empty and not generic, use Description.
    3. Fall back to Device (which equals Value in this case).
  """
  if value != device:
    return value
  if description and not _IsGeneric(description):
    return description
  return device


def _FormatDesignators(parts: str) -> str:
  """Normalise designator list: strip whitespace around commas."""
  return ", ".join(d.strip() for d in parts.split(","))


def ReadFusionBom(path: Path) -> list[dict]:
  """Parse a Fusion 360 BOM CSV, returning rows as dicts."""
  for encoding in ("utf-8-sig", "latin-1", "cp1252"):
    try:
      with open(path, newline="", encoding=encoding) as f:
        reader = csv.DictReader(f)
        return list(reader)
    except UnicodeDecodeError:
      continue
  raise ValueError(f"Could not decode {path} with any supported encoding")


def BuildJlcpcbBomRows(bomRows: list[dict]) -> list[tuple]:
  """Convert Fusion 360 BOM rows to JLCPCB BOM tuples (Comment, Designator, Footprint, Part#)."""
  result = []
  for row in bomRows:
    value = (row.get("Value") or "").strip()
    device = (row.get("Device") or "").strip()
    description = (row.get("Description") or "").strip()
    parts = (row.get("Parts") or "").strip()
    package = (row.get("Package") or "").strip()

    if not parts:
      continue  # skip rows with no designators

    comment = _ChooseComment(value, device, description)
    designator = _FormatDesignators(parts)
    result.append((comment, designator, package, ""))
  return result


# ---------------------------------------------------------------------------
# CPL helpers
# ---------------------------------------------------------------------------

# Fusion 360 CPL has no header; columns are:
#   Designator, X (mm), Y (mm), Rotation, Value, Package
_cplCols = ("Designator", "X", "Y", "Rotation", "Value", "Package")


def ReadFusionCpl(path: Path, layer: str) -> list[tuple]:
  """
  Parse a Fusion 360 CPL CSV.
  Returns list of (Designator, Mid X, Mid Y, Layer, Rotation) tuples.
  """
  rows = []
  with open(path, newline="", encoding="utf-8-sig") as f:
    reader = csv.reader(f)
    for line in reader:
      if not line or not line[0].strip():
        continue
      # Skip any accidental header row (Fusion sometimes adds one)
      if line[0].strip().lower() in ("designator", "refdes", "ref"):
        continue
      if len(line) < 4:
        continue
      designator = line[0].strip()
      midX = f"{float(line[1].strip()):.4f}mm"
      midY = f"{float(line[2].strip()):.4f}mm"
      rotation = float(line[3].strip())
      rows.append((designator, midX, midY, layer, rotation))
  return rows


# ---------------------------------------------------------------------------
# xlsx writing
# ---------------------------------------------------------------------------

_headerFill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_headerFont = Font(bold=True, color="FFFFFF")
_headerAlign = Alignment(horizontal="center", vertical="center")


def _WriteHeader(ws, headers: list[str]) -> None:
  for colIdx, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=colIdx, value=header)
    cell.font = _headerFont
    cell.fill = _headerFill
    cell.alignment = _headerAlign


def _AutofitColumns(ws) -> None:
  for col in ws.columns:
    maxLen = 0
    colLetter = get_column_letter(col[0].column)
    for cell in col:
      if cell.value is not None:
        maxLen = max(maxLen, len(str(cell.value)))
    ws.column_dimensions[colLetter].width = min(maxLen + 4, 60)


def WriteBomXlsx(rows: list[tuple], outPath: Path) -> None:
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "Sheet1"

  headers = ["Comment", "Designator", "Footprint", "JLCPCB Part #（optional）"]
  _WriteHeader(ws, headers)

  for rowIdx, row in enumerate(rows, start=2):
    for colIdx, value in enumerate(row, start=1):
      ws.cell(row=rowIdx, column=colIdx, value=value)

  ws.freeze_panes = "A2"
  _AutofitColumns(ws)
  wb.save(outPath)
  print(f"  BOM written -> {outPath}  ({len(rows)} components)")


def WriteCplXlsx(rows: list[tuple], outPath: Path) -> None:
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "Sheet1"

  headers = ["Designator", "Mid X", "Mid Y", "Layer", "Rotation"]
  _WriteHeader(ws, headers)

  for rowIdx, row in enumerate(rows, start=2):
    for colIdx, value in enumerate(row, start=1):
      ws.cell(row=rowIdx, column=colIdx, value=value)

  ws.freeze_panes = "A2"
  _AutofitColumns(ws)
  wb.save(outPath)
  print(f"  CPL written -> {outPath}  ({len(rows)} placements)")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def ParseArgs() -> argparse.Namespace:
  parser = argparse.ArgumentParser(
    description="Convert Fusion 360 BOM/CPL CSVs to JLCPCB xlsx format.",
    formatter_class=argparse.RawDescriptionHelpFormatter,
    epilog=__doc__,
  )
  parser.add_argument("bomCsv", metavar="BOM.csv", help="Fusion 360 BOM CSV")
  parser.add_argument("frontCplCsv", metavar="FRONT_CPL.csv", help="Fusion 360 front-side CPL CSV")
  parser.add_argument("backCplCsv", metavar="BACK_CPL.csv", nargs="?", default=None,
                      help="Fusion 360 back-side CPL CSV (omit for single-sided boards)")
  parser.add_argument("--bom-out", metavar="PATH", default="bom.xlsx",
                      help="Output path for BOM xlsx (default: ./bom.xlsx)")
  parser.add_argument("--cpl-out", metavar="PATH", default="cpl.xlsx",
                      help="Output path for CPL xlsx (default: ./cpl.xlsx)")
  return parser.parse_args()


def Main() -> None:
  args = ParseArgs()

  bomPath = Path(args.bomCsv)
  frontPath = Path(args.frontCplCsv)
  backPath = Path(args.backCplCsv) if args.backCplCsv else None
  bomOut = Path(args.bom_out)
  cplOut = Path(args.cpl_out)

  for p in filter(None, (bomPath, frontPath, backPath)):
    if not p.exists():
      print(f"Error: file not found: {p}", file=sys.stderr)
      sys.exit(1)

  print("Reading Fusion 360 exports...")
  bomData = ReadFusionBom(bomPath)
  frontCpl = ReadFusionCpl(frontPath, layer="Top")
  backCpl = ReadFusionCpl(backPath, layer="Bottom") if backPath else []

  print(f"  BOM: {len(bomData)} line items")
  print(f"  Front CPL: {len(frontCpl)} placements")
  if backPath:
    print(f"  Back CPL:  {len(backCpl)} placements")
  else:
    print("  Back CPL:  (omitted — single-sided board)")

  bomRows = BuildJlcpcbBomRows(bomData)
  cplRows = frontCpl + backCpl

  print("Writing JLCPCB files...")
  WriteBomXlsx(bomRows, bomOut)
  WriteCplXlsx(cplRows, cplOut)

  print("Done.")


if __name__ == "__main__":
  Main()
